import re
import hashlib
from pathlib import Path
from datetime import datetime, timezone
from typing import Generator

import pandas as pd
from sqlalchemy.orm import Session
from openpyxl import load_workbook

from backend.core.config import CSV_DIR, EXCEL_DIR
from backend.ingestion.indexer import index_document
from backend.ingestion.language_detector import detect_language
from backend.registry.repository import mark_failed, create_source
from backend.registry.models import Source


# ============================================================
# CONFIG
# ============================================================

MAX_TEXT_LENGTH = 200_000
MIN_INDEXABLE_LENGTH = 120


# ============================================================
# FILE HASH
# ============================================================

def compute_file_hash(path: Path) -> str:

    sha = hashlib.sha256()

    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(8192), b""):
            sha.update(chunk)

    return sha.hexdigest()


# ============================================================
# URL EXTRACTION
# ============================================================

URL_PATTERN = re.compile(r"https?://[^\s\"'>)]+", re.IGNORECASE)


def _normalize_url(url: str) -> str:
    return url.strip().rstrip(".,;)\"'")


# ============================================================
# RESOLVE FILE PATH
# ============================================================

def _resolve_path(source: Source) -> Path:

    value = Path(source.source_value).name

    if source.source_type.lower() == "csv":
        return CSV_DIR / value

    return EXCEL_DIR / value


# ============================================================
# STREAM CSV
# ============================================================

def _stream_csv(path: Path) -> Generator[str, None, None]:

    for chunk in pd.read_csv(
        path,
        dtype=str,
        chunksize=400,
        low_memory=True,
        encoding="utf-8",
        on_bad_lines="skip",
    ):

        chunk = chunk.fillna("")

        for _, row in chunk.iterrows():

            values = [
                str(v).strip()
                for v in row.values
                if str(v).strip()
            ]

            if values:
                yield " | ".join(values)


# ============================================================
# STREAM EXCEL USING PANDAS (BEST ENGINE)
# ============================================================

def _stream_excel_pandas(path: Path) -> Generator[str, None, None]:

    try:

        xls = pd.ExcelFile(path)

        for sheet_name in xls.sheet_names:

            df = pd.read_excel(
                xls,
                sheet_name=sheet_name,
                dtype=str
            )

            df = df.fillna("")

            headers = list(df.columns)

            for _, row in df.iterrows():

                row_parts = []

                for col, val in zip(headers, row.values):

                    if not val:
                        continue

                    row_parts.append(f"{col}: {val}")

                if row_parts:
                    yield " | ".join(row_parts)

    except Exception:
        pass


# ============================================================
# STREAM EXCEL USING OPENPYXL (FALLBACK)
# ============================================================

def _stream_excel_openpyxl(path: Path) -> Generator[str, None, None]:

    wb = load_workbook(path, read_only=True, data_only=True)

    try:

        for sheet in wb.worksheets:

            rows = sheet.iter_rows(values_only=True)

            headers = next(rows, None)

            if headers:
                headers = [str(h).strip() if h else f"Column{i}" for i, h in enumerate(headers)]

            for row in rows:

                if not row:
                    continue

                row_parts = []

                for i, value in enumerate(row):

                    if value in (None, ""):
                        continue

                    header = headers[i] if headers and i < len(headers) else f"Column{i}"

                    row_parts.append(f"{header}: {value}")

                if row_parts:
                    yield " | ".join(row_parts)

    finally:
        wb.close()


# ============================================================
# STREAM EXCEL WRAPPER
# ============================================================

def _stream_excel(path: Path) -> Generator[str, None, None]:

    yielded = False

    for row in _stream_excel_pandas(path):

        yielded = True
        yield row

    # fallback if pandas failed
    if not yielded:

        for row in _stream_excel_openpyxl(path):
            yield row


# ============================================================
# MAIN INGEST
# ============================================================

def ingest_tabular(db: Session, source: Source):

    try:

        file_path = _resolve_path(source)

        print("\n📊 INGEST TABULAR:", file_path)

        if not file_path.exists():

            mark_failed(db, source.id, "tabular_file_not_found")

            return {
                "status": "failed",
                "reason": "tabular_file_not_found"
            }

        # ------------------------------------------------------
        # HASH CHECK
        # ------------------------------------------------------

        new_hash = compute_file_hash(file_path)

        if (
            source.file_hash
            and source.file_hash == new_hash
            and source.status == "ingested"
        ):

            return {
                "status": "unchanged",
                "chunks_indexed": 0,
                "urls_discovered": 0,
                "urls_registered": 0,
            }

        # ------------------------------------------------------
        # STREAM CONTENT
        # ------------------------------------------------------

        if source.source_type.lower() == "csv":
            row_stream = _stream_csv(file_path)
        else:
            row_stream = _stream_excel(file_path)

        discovered_urls = set()
        text_parts = []

        row_count = 0
        char_count = 0

        for row_text in row_stream:

            if not row_text:
                continue

            row_count += 1

            text_parts.append(row_text)

            char_count += len(row_text)

            matches = URL_PATTERN.findall(row_text)

            for url in matches:
                discovered_urls.add(_normalize_url(url))

            if char_count > MAX_TEXT_LENGTH:
                break

        full_text = "\n".join(text_parts).strip()

        print("Rows processed:", row_count)
        print("Text length:", len(full_text))

        if len(full_text) < MIN_INDEXABLE_LENGTH:

            mark_failed(db, source.id, "no_indexable_tabular_content")

            return {
                "status": "failed",
                "reason": "no_indexable_tabular_content"
            }

        # ------------------------------------------------------
        # LANGUAGE
        # ------------------------------------------------------

        try:
            language = detect_language(full_text[:5000])
        except Exception:
            language = "unknown"

        # ------------------------------------------------------
        # INDEX DOCUMENT
        # ------------------------------------------------------

        result = index_document(

            source_id=source.id,
            source_type=source.source_type,
            category=source.category,
            language=language,
            source_value=source.source_value,
            page_number=None,
            text=full_text,
        )

        chunks_indexed = result.get("chunks_indexed", 0)

        if chunks_indexed == 0:

            mark_failed(db, source.id, "no_chunks_indexed")

            return {
                "status": "failed",
                "reason": "no_chunks_indexed"
            }

        # ------------------------------------------------------
        # UPDATE SOURCE
        # ------------------------------------------------------

        source.file_hash = new_hash
        source.language = language
        source.last_ingested_at = datetime.now(timezone.utc)
        source.status = "ingested"
        source.error_message = None

        db.commit()

        # ------------------------------------------------------
        # REGISTER DISCOVERED URLS
        # ------------------------------------------------------

        urls_registered = 0

        for url in discovered_urls:

            try:

                new_source = create_source(
                    db=db,
                    source_type="url",
                    source_value=url,
                    origin="discovered",
                    parent_id=source.id,
                    category=source.category,
                )

                if new_source:
                    urls_registered += 1

            except Exception:
                continue

        print("URLs discovered:", len(discovered_urls))
        print("URLs registered:", urls_registered)
        print("Chunks indexed:", chunks_indexed)

        return {

            "status": "ingested",
            "chunks_indexed": chunks_indexed,
            "rows_processed": row_count,
            "urls_discovered": len(discovered_urls),
            "urls_registered": urls_registered,
            "language": language,
            "text_length": len(full_text)
        }

    except Exception as e:

        mark_failed(db, source.id, f"tabular_ingest_error: {e}")

        return {
            "status": "failed",
            "reason": "tabular_ingest_error"
        }